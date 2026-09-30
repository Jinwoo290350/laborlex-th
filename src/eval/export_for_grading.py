"""Phase 2 only (CLAUDE.md §7): run the frozen system on test160 and write an xlsx the
professors can grade: id, question, answer, A1..B3 (blank), เหตุผล.

  make export-160            (requires --confirm-phase2; never used for tuning)
"""

from __future__ import annotations

import argparse
import csv
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from src.eval.judge import CRITERIA

SETS = {"test160": "data/eval/test160.csv", "dev100": "data/eval/dev100.csv"}


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--set", default="test160", choices=SETS)
    ap.add_argument("--confirm-phase2", action="store_true",
                    help="required for test160: config must be frozen (git tag v1.0-eval160)")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--workers", type=int, default=2)
    a = ap.parse_args()
    if a.set == "test160" and not a.confirm_phase2:
        raise SystemExit("test160 is phase-2 only: freeze the config, tag it, then pass --confirm-phase2")

    from src.agent.graph import answer
    with open(SETS[a.set], encoding="utf-8") as f:
        rows = list(csv.DictReader(f))[: a.limit]

    def run(r: dict) -> tuple[dict, str]:
        try:
            return r, answer(r["question"], question_id=r["id"]).markdown
        except Exception as e:  # noqa: BLE001 — keep the row; graders see the failure
            return r, f"[ระบบขัดข้อง: {e!r}]"

    with ThreadPoolExecutor(a.workers) as ex:
        results = list(ex.map(run, rows))

    wb = Workbook()
    ws = wb.active
    ws.title = a.set
    header = ["id", "question", "answer", *CRITERIA, "เหตุผล"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r, md in results:
        ws.append([r["id"], r["question"], md, *[None] * len(CRITERIA), None])
    widths = {"A": 10, "B": 60, "C": 100}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[:3]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    out = Path("docs/results") / f"{datetime.now().astimezone():%Y-%m-%d}_{a.set}_for_grading.xlsx"
    wb.save(out)
    print(f"{len(results)} answers → {out}")


if __name__ == "__main__":
    main()
