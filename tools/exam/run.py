"""CLI for the exam test framework (tools/exam/framework.py). Staging only.

  python -m tools.exam.run intake <file.csv|xlsx> --dataset NAME    validate + draft tagging sheet
  python -m tools.exam.run pilot-dev100 --dataset dev100_pilot      dev100 in intake format (PILOT only)
  python -m tools.exam.run prompts --dataset NAME [--n 3]           build C0/C1 prompts, no API (dry run)
  python -m tools.exam.run run --dataset NAME --cond C0,C1,C2 --ids a,b --dry-run   fetch + size, no Gemini
  python -m tools.exam.run run --dataset NAME --cond C0,C1 --ids a,b --live
                                                                    Gemini; needs EXAM_ALLOW_GEMINI=1, ≤ 12 questions
  python -m tools.exam.run scope --dataset NAME                    count in/out of scope + buckets (before any run)
  python -m tools.exam.run review-sheet --dataset NAME              blind xlsx for the human reviewer
  python -m tools.exam.run report --dataset NAME [--reviews file.xlsx]

Files: data/eval/exams/<dataset>/{questions.csv, tags.csv, tags.xlsx, prompts/, answers/, review*.xlsx,
review_key.json, report.md}. data/ is gitignored.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import random
import sys
from pathlib import Path

from pydantic import BaseModel

sys.path.insert(0, ".")
from tools.exam import framework as fw

ROOT = Path("data/eval/exams")
MAX_LIVE = 12


class ExamAnswer(BaseModel):
    issues: list[str]
    legal_rules: list[str]
    application: str
    calculation: str = ""
    conclusion: str


def _dir(ds: str) -> Path:
    return ROOT / ds


def _write_csv(path: Path, rows: list[dict], cols: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def _xlsx(path: Path, rows: list[dict], cols: list[str], lists: dict[str, list[str]] | None = None) -> None:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    for r in rows:
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", str(r.get(c, ""))) for c in cols])
    for name, opts in (lists or {}).items():
        col = cols.index(name) + 1
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = ws.cell(1, col).column_letter
        dv.add(f"{letter}2:{letter}{len(rows) + 1}")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


def load(ds: str) -> tuple[list[dict], dict[str, dict]]:
    qs = fw.read_table(_dir(ds) / "questions.csv")
    tag_path = _dir(ds) / "tags.xlsx"           # the person edits the xlsx; csv is the untouched draft
    tags = fw.read_table(tag_path if tag_path.exists() else _dir(ds) / "tags.csv")
    return qs, {t["exam_id"]: t for t in tags}


def cmd_intake(a) -> None:
    rows = fw.read_table(Path(a.file))
    if errs := fw.validate_intake(rows):
        sys.exit("intake errors:\n  " + "\n  ".join(errs))
    tags = fw.tag_rows(rows)
    _write_csv(_dir(a.dataset) / "questions.csv", rows, fw.INTAKE_COLS)
    _write_csv(_dir(a.dataset) / "tags.csv", tags, fw.TAG_COLS)
    _xlsx(_dir(a.dataset) / "tags.xlsx", tags, fw.TAG_COLS,
          {"calculation": ["Y", "N"], "needs_ruling": ["Y", "N"], "in_scope": ["yes", "review", "no"]})
    cnt = {b: sum(t["bucket"] == b for t in tags) for b in ("0", "1", "2", ">2")}
    print(f"{len(rows)} questions → {_dir(a.dataset)}  draft buckets {cnt}")
    print(fw.scope_summary(tags))
    print("next: confirm sections / tags in tags.xlsx (sections_confirmed, n_issues, calculation, needs_ruling)")


def cmd_pilot_dev100(a) -> None:
    with open("data/eval/dev100.csv", encoding="utf-8") as fh:
        dev = list(csv.DictReader(fh))
    rows = [{"exam_id": r["id"], "source": "dev100_pilot", "exam_year": "", "exam_ref": "dev100",
             "question_no": r["id"], "question": r["question"], "model_answer": r["gold_answer"],
             "answer_source": "client sheet column H", "event_date": r.get("event_date", ""),
             "notes": "PILOT ONLY — not a Ram/Bar exam question"} for r in dev]
    tmp = _dir(a.dataset) / "_intake.csv"
    _write_csv(tmp, rows, fw.INTAKE_COLS)
    a.file = str(tmp)
    cmd_intake(a)


def cmd_prompts(a) -> None:
    qs, tags = load(a.dataset)
    out = _dir(a.dataset) / "prompts"
    out.mkdir(parents=True, exist_ok=True)
    for q in qs[: a.n]:
        secs = fw.confirmed_sections(tags[q["exam_id"]])
        for cond in ("C0", "C1"):
            version, text = fw.build_prompt(cond, q["question"], secs)
            (out / f"{q['exam_id']}_{cond}.txt").write_text(f"# prompt exam_{cond.lower()} v{version}\n{text}",
                                                             encoding="utf-8")
            print(f"{q['exam_id']} {cond} v{version}: {len(text):,} chars  sections {secs if cond == 'C1' else '-'}")
    print(f"dry run only, no API call → {out}")


def _price() -> tuple[float, float]:
    from src.config import settings
    return settings.gemini_price_in, settings.gemini_price_out


_SRC_CACHE: dict[str, list[dict]] = {}


def source_text(q: dict) -> tuple[str, str]:
    """(question, model_answer). Scoped web datasets keep no text in the repo: fetch the source page at run
    time (in memory only, ≥ 1.2 s between requests) and parse it with tools/exam/parse_mjsheet.py."""
    if q.get("question"):
        return q["question"], q["model_answer"]
    url = q.get("source_url") or ""
    if not url:
        raise SystemExit(f"{q['exam_id']}: no text and no source_url")
    if url not in _SRC_CACHE:
        import time
        import urllib.request

        from tools.exam.parse_mjsheet import parse_page
        req = urllib.request.Request(urllib.parse.quote(url, safe=":/"), headers={"User-Agent": "laborlex-th-research"})
        html = urllib.request.urlopen(req, timeout=30).read().decode("utf-8", "ignore")
        _SRC_CACHE[url] = parse_page(html, url)
        time.sleep(1.2)
    row = next((r for r in _SRC_CACHE[url] if str(r["question_no"]) == str(q["question_no"])), None)
    if not row or "NO_ANSWER" in row["parse_flags"]:
        raise SystemExit(f"{q['exam_id']}: question {q['question_no']} not found / no model answer at {url}")
    return row["question"], row["model_answer"]


def _sections_for(tag: dict) -> list[str]:
    """Scoped datasets: only human-confirmed sections (never the unreviewed draft)."""
    if tag.get("in_scope"):
        confirmed = [x.strip() for x in (tag.get("sections_confirmed") or "").split(";") if x.strip()]
        if not confirmed:
            raise SystemExit(f"{tag['exam_id']}: sections_confirmed is empty — review it before running")
        bad = [x for x in confirmed if x.split(":")[0] not in fw.SCOPE_LAWS]
        if bad:
            raise SystemExit(f"{tag['exam_id']}: confirmed sections outside LPA/LRA scope: {bad}")
        return confirmed
    return fw.confirmed_sections(tag)


def _env_record() -> dict:
    import subprocess
    head = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True, check=False).stdout.strip()
    dirty = subprocess.run(["git", "diff", "--quiet", "HEAD", "--", "src"], check=False).returncode != 0
    from src.config import settings
    return {"git_head": head, "src_modified_vs_head": dirty, "model": settings.gemini_model}


def _staging_record(label: str) -> dict:
    """What an explicitly allowed staging C2 run ran on: branch, changed src/ and data files, hash of their diff."""
    import hashlib
    import subprocess
    def git(*args: str) -> str:
        return subprocess.run(["git", *args], capture_output=True, text=True, check=False).stdout
    paths = ["src", "config", "data/processed/issues.yaml", "data/processed/dev100_issues.json"]
    files = sorted(set(git("diff", "--name-only", "HEAD", "--", *paths).split())
                   | set(git("ls-files", "--others", "--exclude-standard", "--", *paths).split())
                   | {f for f in paths[2:] if Path(f).exists()})      # data files may be gitignored
    h = hashlib.sha256()
    for f in files:
        h.update(f.encode() + Path(f).read_bytes() if Path(f).exists() else f.encode())
    return {"staging_label": label, "branch": git("rev-parse", "--abbrev-ref", "HEAD").strip(),
            "git_head": git("rev-parse", "--short", "HEAD").strip(), "files": files,
            "files_sha256": h.hexdigest()[:16]}


def cmd_run(a) -> None:
    import time
    conds = a.cond.split(",")
    if not a.dry_run and not a.live:
        sys.exit("refusing: Gemini calls need --live (or use --dry-run)")
    if a.live and os.environ.get("EXAM_ALLOW_GEMINI") != "1":
        sys.exit("refusing: set EXAM_ALLOW_GEMINI=1 after the dataset and budget are approved")
    qs, tags = load(a.dataset)
    ids = a.ids.split(",") if a.ids else []
    todo = [q for q in qs if q["exam_id"] in ids]
    blocked = [q["exam_id"] for q in todo if not fw.runnable(tags[q["exam_id"]])]
    if blocked:
        sys.exit(f"refusing: not in scope (in_scope must be 'yes' in tags.xlsx): {blocked}")
    if not todo or len(todo) > MAX_LIVE:
        sys.exit(f"refusing: pass --ids with 1–{MAX_LIVE} exam_ids (got {len(todo)})")
    env = _env_record()
    if "C2" in conds and env["src_modified_vs_head"]:
        if not a.allow_src_diff:
            sys.exit("refusing: src/ differs from HEAD — C2 must run the current committed retrieval unchanged "
                     "(staging fixes: --allow-src-diff LABEL records the diff and writes *_C2-LABEL.json)")
        env |= _staging_record(a.allow_src_diff)
    if a.dry_run:                       # fetch text + build C0/C1 prompts, report sizes; no Gemini, nothing written
        pin, pout = _price()
        for q in todo:
            secs = _sections_for(tags[q["exam_id"]])
            question, model_answer = source_text(q)
            sizes = {c: len(fw.build_prompt(c, question, secs)[1]) for c in ("C0", "C1")}
            print(f"{q['exam_id']:24} bucket {fw.bucket(len(secs)):>2} C1 sections {';'.join(secs):45} "
                  f"q {len(question):>5} ch · model answer {len(model_answer):>5} ch · prompt C0 {sizes['C0']:>5} ch "
                  f"C1 {sizes['C1']:>6} ch")
        print(f"dry run: no Gemini call, no text written · env {env} · price in ${pin}/M out ${pout}/M")
        return
    print(fw.scope_summary(list(tags.values())))
    from src.llm import USAGE, generate_json
    pin, pout = _price()
    out = _dir(a.dataset) / "answers"
    out.mkdir(parents=True, exist_ok=True)
    for q in todo:
        secs = _sections_for(tags[q["exam_id"]])
        question, model_answer = source_text(q)
        for cond in conds:
            n0, t0, err = len(USAGE.log), time.monotonic(), None
            ans: dict = {}
            try:
                if cond == "C2":
                    from src.agent.graph import answer
                    s = answer(question)
                    keys = sorted({k for t in ([s.answer.model_dump()] if s.answer else []) for k in _keys(t)})
                    from src.agent.status import effective_status
                    ans = {"markdown": s.markdown, "cited_keys": keys, "c2_state": _c2_state(s),
                           "status": effective_status(s), "review_reasons": s.degraded}
                    if ans["status"] != "ok":      # #22: never scored as a normal answer
                        err = f"INVALID_{ans['status'].upper()}"
                else:
                    version, text = fw.build_prompt(cond, question, secs)
                    ans = generate_json(text, ExamAnswer, name=f"exam_{cond}").model_dump() | {"prompt_version": version}
            except Exception as e:  # noqa: BLE001 — record the failure and continue
                err = repr(e)
            calls = USAGE.log[n0:]
            tin, tout = sum(c["in"] for c in calls if not c["cached"]), sum(c["out"] for c in calls if not c["cached"])
            text_out = fw.answer_text(ans) if ans else ""
            chk = fw.auto_checks(secs, text_out, model_answer)
            rec = {"exam_id": q["exam_id"], "condition": cond, "bucket": fw.bucket(len(secs)),
                   "sections_confirmed": secs, "answer": ans, "checks": chk,
                   "conclusion": ans.get("conclusion", "") if cond != "C2" else "",
                   "calculation": ans.get("calculation", "") if cond != "C2" else "",
                   "cited_sections": chk["cited"], "latency_s": round(time.monotonic() - t0, 1),
                   "tokens": {"in": tin, "out": tout, "calls": len(calls)},
                   "cost_usd": round((tin * pin + tout * pout) / 1e6, 4),
                   "failure": err or ("EMPTY_ANSWER" if not text_out else None), "env": env}
            tag = f"{cond}-{a.allow_src_diff}" if cond == "C2" and env.get("staging_label") else cond
            (out / f"{q['exam_id']}_{tag}.json").write_text(json.dumps(rec, ensure_ascii=False, indent=1, default=str),
                                                             encoding="utf-8")
            print(f"{q['exam_id']} {cond}: recall {chk['section_recall']} missing {chk['missing']} "
                  f"{rec['latency_s']}s ${rec['cost_usd']} {rec['failure'] or ''}")
    print(f"Gemini calls {USAGE.calls} (cached {USAGE.cached}) tokens in {USAGE.input_tokens} out {USAGE.output_tokens}")


def _c2_state(s) -> dict:
    """What C2 retrieved and decided, per issue (recording only; the pipeline itself is unchanged)."""
    def ck(d: dict) -> str:
        return d.get("citation_key") or d.get("deka_no") or d.get("key") or d.get("case_key") or str(d.get("id"))
    return {
        "facts_asked": (s.facts.asked if s.facts else []),
        "issues": [{"code": i.code, "reason": i.reason} for i in s.issues],
        "candidates": {k: [ck(d) for d in v][:20] for k, v in s.candidates.items()},
        "cases": {k: [{"case": ck(d), "court": d.get("court"), "score": d.get("score")} for d in v] for k, v in s.cases.items()},
        "selected": s.selected,
        "elements": {k: [e.model_dump() for e in v] for k, v in s.elements.items()},
        "calcs": {k: (v.model_dump() if hasattr(v, "model_dump") else str(v)) for k, v in s.calcs.items()},
        "removed_citations": s.removed_citations,
        "status": s.status, "degraded": s.degraded,
        "trace": [{k: t.get(k) for k in ("node", "latency_ms", "in_tokens", "out_tokens", "error")} for t in s.trace],
    }


def _keys(obj) -> list[str]:
    """Citation keys anywhere in a C2 AnswerJSON (fields ending in _keys / citation_key)."""
    out: list[str] = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            if (k.endswith("_keys") or k == "citation_key") and v:
                out += [x for x in (v if isinstance(v, list) else [v]) if isinstance(x, str)]
            else:
                out += _keys(v)
    elif isinstance(obj, list):
        for v in obj:
            out += _keys(v)
    return out


def _answers(ds: str) -> dict:
    return {(j["exam_id"], j["condition"]): j for p in (_dir(ds) / "answers").glob("*.json")
            for j in [json.loads(p.read_text(encoding="utf-8"))]}


def cmd_scope(a) -> None:
    _, tags = load(a.dataset)
    print(fw.scope_summary(list(tags.values())))


def cmd_review_sheet(a) -> None:
    qs, _ = load(a.dataset)
    qmap = {q["exam_id"]: q for q in qs}
    ans = _answers(a.dataset)
    rows = []
    for (eid, cond), j in ans.items():
        rows.append({"review_id": "", "exam_id": eid,
                     "question": qmap[eid]["question"] or "(see source_url)",
                     "model_answer": qmap[eid]["model_answer"] or "(see source_url)",
                     "source_url": qmap[eid].get("source_url", ""), "ai_answer": fw.answer_text(j["answer"]),
                     "missing_sections_auto": ";".join(j["checks"]["missing"]), "_cond": cond})
    random.Random(20261005).shuffle(rows)
    for k, r in enumerate(rows, 1):
        r["review_id"] = f"R{k:03d}"
    cols = ["review_id", "exam_id", "source_url", "question", "model_answer", "ai_answer", "missing_sections_auto",
            *fw.REVIEW_LABELS, "main_cause", "note"]
    _xlsx(_dir(a.dataset) / "review.xlsx", rows, cols, {**fw.REVIEW_LABELS, "main_cause": list(fw.CAUSES)})
    (_dir(a.dataset) / "review_key.json").write_text(json.dumps({r["review_id"]: r["_cond"] for r in rows},
                                                                indent=1), encoding="utf-8")
    print(f"{len(rows)} rows → {_dir(a.dataset) / 'review.xlsx'} (condition hidden; key in review_key.json)")


def cmd_report(a) -> None:
    _, tags = load(a.dataset)
    answers = {k: v["checks"] for k, v in _answers(a.dataset).items()}
    reviews = {}
    if a.reviews:
        key = json.loads((_dir(a.dataset) / "review_key.json").read_text(encoding="utf-8"))
        reviews = {(r["exam_id"], key[r["review_id"]]): r for r in fw.read_table(Path(a.reviews))}
    text = fw.report(list(tags.values()), answers, reviews, a.dataset)
    (_dir(a.dataset) / "report.md").write_text(text, encoding="utf-8")
    print(text)


def main() -> None:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("intake")
    p.add_argument("file")
    p.add_argument("--dataset", required=True)
    sub.add_parser("pilot-dev100").add_argument("--dataset", default="dev100_pilot")
    p = sub.add_parser("prompts")
    p.add_argument("--dataset", required=True)
    p.add_argument("--n", type=int, default=3)
    p = sub.add_parser("run")
    p.add_argument("--dataset", required=True)
    p.add_argument("--cond", default="C0,C1")
    p.add_argument("--ids", default="")
    p.add_argument("--live", action="store_true")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--allow-src-diff", default="", metavar="LABEL",
                   help="run C2 on uncommitted staging src/ changes; recorded in env, output *_C2-LABEL.json")
    sub.add_parser("scope").add_argument("--dataset", required=True)
    sub.add_parser("review-sheet").add_argument("--dataset", required=True)
    p = sub.add_parser("report")
    p.add_argument("--dataset", required=True)
    p.add_argument("--reviews", default="")
    a = ap.parse_args()
    {"intake": cmd_intake, "pilot-dev100": cmd_pilot_dev100, "prompts": cmd_prompts, "run": cmd_run,
     "review-sheet": cmd_review_sheet, "scope": cmd_scope, "report": cmd_report}[a.cmd](a)


if __name__ == "__main__":
    main()
