"""Exam test framework (2026-10-05): real Ramkhamhaeng / Thai Bar labour-law exam questions, split by the
number of sections the model answer needs (1 / 2 / >2), answered under three conditions, diagnosed by
stage. Plan: docs/design/2026-10-05_exam_test_plan.md. Staging only; nothing here touches production.

Conditions
  C0  Gemini, question only                        (what it already knows)
  C1  Gemini, question + text of the model answer's sections from our DB ("teaching the law")
  C2  our system (src.agent.graph.answer), finds the law itself
Gemini is called only by `run --live` (see run.py); everything else is local.

Dataset honesty: every file carries `source`. Only REAL_SOURCES count as exam results; any other
source (e.g. dev100_pilot) is stamped "PILOT — not Ram/Bar results" in every report.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from tools.exam.statutes import statutes, unknown

REAL_SOURCES = {"ram", "bar"}
INTAKE_COLS = ["exam_id", "source", "exam_year", "exam_ref", "question_no", "question", "model_answer",
               "answer_source", "event_date", "notes", "source_url", "sections_needed_draft", "sections_supporting_draft",
               "in_scope", "scope_reason"]
TAG_COLS = ["exam_id", "source", "exam_year", "in_scope", "scope_reason", "auto_sections", "sections_confirmed",
            "n_sections", "bucket", "n_issues", "calculation", "cross_act", "needs_ruling", "old_law_risk",
            "tagger_note"]
# Benchmark scope (2026-10-05): the laws the system supports. Buckets count only these laws' needed sections.
SCOPE_LAWS = ("LPA2541", "LRA2518")
CONDITIONS = ("C0", "C1", "C2")
# Human review: one row per (exam, condition), condition hidden from the reviewer.
REVIEW_LABELS = {
    "conclusion": ["ถูก", "ถูกบางส่วน", "ผิด"],                       # A2 / B3 — headline
    "issue_spotting": ["ครบ", "ขาด", "เกิน", "ขาดและเกิน"],            # B0
    "legal_correctness": ["ถูก", "ผิดบางส่วน", "ผิด"],                 # B1 / A2: right section, right rule
    "fact_application": ["ถูก", "ผิดบางส่วน", "ผิด", "ไม่ปรับบท"],      # B2
}
CAUSES = {  # one main cause per wrong / partly wrong answer
    "issue_spotting": "ไม่เห็นประเด็น / เพิ่มประเด็นเกิน (B0)",
    "section_selection": "ไม่ได้ใช้มาตราที่ต้องใช้ หรือหยิบมาตราผิด (B1)",
    "wrong_version": "ใช้กฎหมายผิดฉบับ / ผิดช่วงเวลา (B1/A2)",
    "legal_principle": "มาตราถูกแต่เข้าใจหลักผิด (B1/A2)",
    "fact_application": "ปรับบทกับข้อเท็จจริงผิด (B2)",
    "calculation": "คำนวณผิด (A2)",
    "conclusion": "สรุปผิดหรือขัดกับการวิเคราะห์ (B3/A2)",
    "other": "อื่น ๆ (ระบุในหมายเหตุ)",
}


# ---------- intake & bucketing ----------

def section_set(text: str) -> list[str]:
    """Section-level keys named in a text: วรรค / อนุมาตรา collapse into the section; 17 and 17/1 stay two.
    Law per mention: tools/exam/statutes.py (not the dev100 gold importer)."""
    return statutes(text or "")


def bucket(n: int) -> str:
    return "0" if n == 0 else "1" if n == 1 else "2" if n == 2 else ">2"


def read_table(path: Path) -> list[dict]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True).active
        rows = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h else "" for h in next(rows)]
        return [{h: ("" if v is None else str(v).strip()) for h, v in zip(head, r)} for r in rows if any(r)]
    with open(path, encoding="utf-8-sig") as fh:
        return [{k.strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(fh)]


def validate_intake(rows: list[dict]) -> list[str]:
    errs = []
    if rows and (missing := [c for c in ("exam_id", "source", "question", "model_answer") if c not in rows[0]]):
        return [f"missing columns: {missing}"]
    ids = Counter(r["exam_id"] for r in rows)
    errs += [f"duplicate exam_id {i}" for i, n in ids.items() if n > 1]
    for r in rows:
        if not r["question"]:
            errs.append(f"{r['exam_id']}: empty question")
        if not r["model_answer"]:
            errs.append(f"{r['exam_id']}: no model answer (cannot bucket or score)")
    return errs


def amendment_years() -> dict[str, list[int]]:
    """LPA section → years of amending acts (staging table from the krisdika consolidated PDF)."""
    p = Path("data/staging/kb/lpa_amendments.jsonl")
    out: dict[str, list[int]] = defaultdict(list)
    if p.exists():
        for line in p.read_text(encoding="utf-8").splitlines():
            r = json.loads(line)
            y = re.search(r"พ\.ศ\.\s*(25\d\d)", r.get("amending_instrument") or "")
            if r.get("section") and y:
                out[f"LPA2541:{r['section']}"].append(int(y.group(1)))
    return out


def tag_rows(rows: list[dict]) -> list[dict]:
    """Draft tagging sheet: sections parsed from the model answer, for a person to confirm."""
    amend = amendment_years()
    out = []
    for r in rows:
        # Prefer the parser's "needed for the decision" draft; counting every section in a model answer overcounts.
        needed = [s for s in (r.get("sections_needed_draft") or "").split(";") if s]
        secs = needed or section_set(r["model_answer"])
        if r.get("in_scope"):                     # scoped datasets (Ram): bucket on LPA/LRA needed sections only
            secs = [s for s in secs if s.split(":")[0] in SCOPE_LAWS]
        year = int(r["exam_year"]) if (r.get("exam_year") or "").isdigit() else None
        risk = sorted({s for s in secs if year and any(y > year for y in amend.get(s, []))})
        out.append({"exam_id": r["exam_id"], "source": r["source"], "exam_year": r.get("exam_year", ""),
                    "in_scope": r.get("in_scope", ""), "scope_reason": r.get("scope_reason", ""),
                    "auto_sections": ";".join(secs), "sections_confirmed": "", "n_sections": len(secs),
                    "bucket": bucket(len(secs)), "n_issues": "", "calculation": "",
                    "cross_act": "Y" if len({s.split(":")[0] for s in secs}) > 1 else "N",
                    "needs_ruling": "", "old_law_risk": ";".join(risk), "tagger_note": ""})
    return out


def runnable(tag: dict) -> bool:
    """Unscoped datasets run as before; scoped ones only when a person has settled in_scope = yes."""
    return tag.get("in_scope", "") in ("", "yes")


def scope_summary(tags: list[dict]) -> str:
    cnt = Counter(t.get("in_scope") or "unscoped" for t in tags)
    keep = [t for t in tags if runnable(t)]
    by_b = Counter(bucket(len(confirmed_sections(t))) for t in keep)
    buckets = f"1: {by_b['1']} · 2: {by_b['2']} · >2: {by_b['>2']} · 0: {by_b['0']}"
    lines = [f"questions {len(tags)} · in_scope " + " · ".join(f"{k} {v}" for k, v in sorted(cnt.items())),
             f"runnable (in_scope = yes) {len(keep)} → buckets {buckets}"]
    reasons = Counter(t.get("scope_reason") for t in tags if t.get("in_scope") in ("no", "review"))
    lines += [f"  {n} × {r}" for r, n in reasons.most_common()]
    return "\n".join(lines)


def confirmed_sections(tag: dict) -> list[str]:
    raw = tag.get("sections_confirmed") or tag.get("auto_sections") or ""
    return [s.strip() for s in raw.split(";") if s.strip()]


# ---------- conditions ----------

def law_text(sections: list[str], on=None) -> str:
    """C1 law block from the DB (paragraphs in force on `on`; None = current). Never from memory."""
    from src.index.tools import _in_force, get_section
    blocks = []
    for key in sections:
        law, sec = key.split(":", 1)
        paras = [p for p in get_section(law, sec) if _in_force(p, on)]
        body = "\n".join(p["text"] for p in paras) if paras else "(ไม่พบตัวบทในฐานข้อมูล)"
        blocks.append(f"[{key}]\n{body}")
    return "\n\n".join(blocks)


def build_prompt(cond: str, question: str, sections: list[str], on=None) -> tuple[str, str]:
    from src.agent.nodes.common import prompt
    version, body = prompt("exam_c0" if cond == "C0" else "exam_c1")
    text = body.replace("{question}", question)
    if cond == "C1":
        text = text.replace("{law_text}", law_text(sections, on))
    return version, text


def answer_text(ans: dict) -> str:
    """Flatten a C0/C1 JSON answer or a C2 record into one text for the automatic checks."""
    if ans.get("markdown"):
        return ans["markdown"]
    parts = [*ans.get("issues", []), *ans.get("legal_rules", []), ans.get("application", ""),
             ans.get("calculation", ""), ans.get("conclusion", "")]
    return "\n".join(p for p in parts if p)


# ---------- automatic checks (no API) ----------

_AMOUNT = re.compile(r"(\d[\d,]*(?:\.\d+)?)\s*บาท")


def amounts(text: str) -> set[str]:
    """Baht amounts, normalised: "12,000.00 บาท" → "12000"."""
    return {re.sub(r"\.0+$", "", a.replace(",", "")) for a in _AMOUNT.findall(text or "")}


def auto_checks(required: list[str], answer: str, model_answer: str) -> dict:
    named = section_set(answer)
    bad = set(unknown(named))                 # sections the named law does not have: reported, never counted
    cited = set(named) - bad
    req = set(required)
    gold_amt = amounts(model_answer)
    return {"required": sorted(req), "cited": sorted(cited), "unknown_cited": sorted(bad),
            "missing": sorted(req - cited),
            "extra": sorted(cited - req), "all_sections": not (req - cited),
            "section_recall": round(len(req & cited) / len(req), 3) if req else None,
            "amounts_expected": sorted(gold_amt),
            "amounts_match": (gold_amt <= amounts(answer)) if gold_amt else None}


# ---------- diagnosis & report ----------

def correct(review: dict | None) -> bool | None:
    if not review or not review.get("conclusion"):
        return None
    return review["conclusion"] == "ถูก"


def diagnose(c0: bool | None, c1: bool | None, c2: bool | None, c2_all_sections: bool | None) -> str:
    """Where the failure sits, from the human conclusion verdict under each condition."""
    if None in (c0, c1, c2):
        return "incomplete"
    if c0 and c1 and c2:
        return "solved"
    if not c1:
        return "reasoning_gap"            # wrong even with the right law in front of it
    if not c2:
        return "retrieval_gap" if c2_all_sections is False else "system_reasoning_gap"
    return "knowledge_gap_fixed" if not c0 else "solved_c1_c2"


DIAGNOSIS_TH = {
    "solved": "ทุกเงื่อนไขตอบถูก",
    "knowledge_gap_fixed": "ไม่รู้กฎหมาย แต่สอนตัวบทแล้วตอบถูก (C0 ผิด → C1, C2 ถูก)",
    "solved_c1_c2": "C0 ผิดเฉพาะบางเคส อื่นถูก",
    "reasoning_gap": "มีตัวบทที่ถูกแล้วยังตอบผิด (ปัญหาการให้เหตุผล/ปรับบท)",
    "retrieval_gap": "สอนตัวบทแล้วถูก แต่ระบบหามาตราไม่ครบ (ปัญหาการค้นตัวบท)",
    "system_reasoning_gap": "สอนตัวบทแล้วถูก ระบบได้มาตราครบแต่ยังตอบผิด",
    "incomplete": "ยังตรวจไม่ครบ",
}


def report(tags: list[dict], answers: dict, reviews: dict, dataset: str) -> str:
    """answers[(exam_id, cond)] = auto_checks dict; reviews[(exam_id, cond)] = human review row."""
    real = {t["source"] for t in tags} <= REAL_SOURCES and bool(tags)
    out = [f"# Exam test report — dataset `{dataset}`", ""]
    if not real:
        srcs = sorted({t["source"] for t in tags})
        banner = (f"> ⚠️ **PILOT — ไม่ใช่ผลข้อสอบราม/เนติ** (sources: {srcs}) ใช้ทดสอบวิธีเท่านั้น"
                  " ห้ามนำตัวเลขไปนำเสนอเป็นผลข้อสอบจริง")
        out += [banner, ""]
    out += ["## Scope", "", "```", scope_summary(tags), "```", ""]
    tags = [t for t in tags if runnable(t)]
    by_b = defaultdict(list)
    for t in tags:
        by_b[bucket(len(confirmed_sections(t)))].append(t["exam_id"])
    out += ["## ตาราง 3×3: ข้อสรุปถูก (คนตรวจ) · อ้างมาตราครบ (อัตโนมัติ)", "",
            "| กลุ่ม | n | " + " | ".join(CONDITIONS) + " |", "|---|---:|" + "---|" * len(CONDITIONS)]
    for b in ("1", "2", ">2"):
        ids = by_b.get(b, [])
        cells = []
        for c in CONDITIONS:
            rv = [correct(reviews.get((i, c))) for i in ids]
            done = [x for x in rv if x is not None]
            sec = [answers[(i, c)]["all_sections"] for i in ids if (i, c) in answers]
            cells.append(f"{sum(done)}/{len(done)} ถูก · {sum(sec)}/{len(sec)} มาตราครบ" if ids else "—")
        out.append(f"| {b} มาตรา | {len(ids)} | " + " | ".join(cells) + " |")
    diag = {i: diagnose(*(correct(reviews.get((i, c))) for c in CONDITIONS),
                        answers.get((i, "C2"), {}).get("all_sections")) for t in tags for i in [t["exam_id"]]}
    out += ["", "## ผิดที่ขั้นไหน (C0/C1/C2)", "", "| กลุ่ม | " + " | ".join(DIAGNOSIS_TH) + " |",
            "|---|" + "---:|" * len(DIAGNOSIS_TH)]
    for b in ("1", "2", ">2"):
        cnt = Counter(diag[i] for i in by_b.get(b, []))
        out.append(f"| {b} มาตรา | " + " | ".join(str(cnt[k]) for k in DIAGNOSIS_TH) + " |")
    out += ["", "หัวตาราง: " + " · ".join(f"`{k}` = {v}" for k, v in DIAGNOSIS_TH.items()), "",
            "## สาเหตุหลักของคำตอบที่ผิด (คนตรวจ)", "", "| กลุ่ม | เงื่อนไข | " + " | ".join(CAUSES) + " |",
            "|---|---|" + "---:|" * len(CAUSES)]
    for b in ("1", "2", ">2"):
        for c in CONDITIONS:
            cnt = Counter((reviews.get((i, c)) or {}).get("main_cause") for i in by_b.get(b, []))
            out.append(f"| {b} | {c} | " + " | ".join(str(cnt[k]) for k in CAUSES) + " |")
    out += ["", "สาเหตุ: " + " · ".join(f"`{k}` = {v}" for k, v in CAUSES.items()),
            "", "Judge PASS ไม่ได้ใช้เป็นตัวเลขหลัก (ตรงกับทนายเพียง 58% บน dev100)"]
    return "\n".join(out)
